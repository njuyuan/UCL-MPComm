"""Generate llm-trends-network-requirements.pptx next to this script.

Usage:
    pip install python-pptx
    python3 make_ppt.py
"""

import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

FONT = "Microsoft YaHei"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLUE = RGBColor(0x2E, 0x75, 0xB6)
LIGHT_BLUE = RGBColor(0xDE, 0xEB, 0xF7)
ORANGE = RGBColor(0xED, 0x7D, 0x31)
GREEN = RGBColor(0x3D, 0x8B, 0x37)
PURPLE = RGBColor(0x70, 0x30, 0xA0)
GRAY = RGBColor(0x7F, 0x7F, 0x7F)
LIGHT_GRAY = RGBColor(0xF4, 0xF6, 0xF9)
MID_GRAY = RGBColor(0xD0, 0xD5, 0xDD)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x26, 0x26, 0x26)
GPU_FILL = RGBColor(0x9D, 0xC3, 0xE6)
TEAL = RGBColor(0x1A, 0x9E, 0x9E)


def set_font(run, size, bold=False, color=DARK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", FONT)


def write(tf, lines, size=11, color=DARK, align=None, bold=False, spacing=None):
    """lines: list of str or list of (text, bold, color) segment lists."""
    tf.word_wrap = True
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if align is not None:
            p.alignment = align
        if spacing is not None:
            p.space_after = Pt(spacing)
        segs = line if isinstance(line, list) else [(line, bold, color)]
        for text, b, c in segs:
            r = p.add_run()
            r.text = text
            set_font(r, size, b, c)


def box(slide, x, y, w, h, text="", fill=WHITE, line=MID_GRAY, size=10,
        bold=False, color=DARK, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, line_w=1.0, dash=False):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.12
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = 4
    s.shadow.inherit = False
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = anchor
    if text:
        write(tf, text.split("\n") if isinstance(text, str) else text,
              size=size, color=color, align=align, bold=bold)
    return s


def label(slide, x, y, w, h, text, size=9, color=DARK, bold=False,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.0)
    tf.vertical_anchor = anchor
    write(tf, text.split("\n") if isinstance(text, str) else text,
          size=size, color=color, align=align, bold=bold)
    return tb


def arrow(slide, x1, y1, x2, y2, color=GRAY, width=1.5, dash=False,
          head=True, tail=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = 4
    ln = c.line._get_or_add_ln()
    if tail:
        e = etree.SubElement(ln, qn("a:headEnd"))
        e.set("type", "triangle")
        e.set("w", "med")
        e.set("len", "med")
    if head:
        e = etree.SubElement(ln, qn("a:tailEnd"))
        e.set("type", "triangle")
        e.set("w", "med")
        e.set("len", "med")
    return c


def gpu_grid(slide, x, y, cols, rows, size=0.32, gap=0.1, fill=GPU_FILL, text="GPU"):
    if size < 0.3:
        text = ""
    for r in range(rows):
        for c in range(cols):
            box(slide, x + c * (size + gap), y + r * (size + gap), size, size,
                text, fill=fill, line=BLUE, size=6, shape=MSO_SHAPE.RECTANGLE,
                line_w=0.75)


# ---------------------------------------------------------------- page frame

def frame(prs, idx, title, key_msg, sections, notes, footer=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    box(slide, 0, 0, 13.333, 0.12, fill=NAVY, line=None, shape=MSO_SHAPE.RECTANGLE)
    label(slide, 0.4, 0.25, 12.5, 0.65, title, size=24, bold=True, color=NAVY,
          align=PP_ALIGN.LEFT)

    km = box(slide, 0.4, 0.98, 12.53, 0.62, fill=LIGHT_BLUE, line=None,
             shape=MSO_SHAPE.RECTANGLE, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT)
    km.text_frame.margin_left = Inches(0.2)
    write(km.text_frame, [[("核心判断：", True, NAVY), (key_msg, False, DARK)]], size=13)
    box(slide, 0.4, 0.98, 0.07, 0.62, fill=BLUE, line=None, shape=MSO_SHAPE.RECTANGLE)

    panel = box(slide, 0.4, 1.78, 6.3, 5.12, fill=LIGHT_GRAY, line=MID_GRAY,
                shape=MSO_SHAPE.RECTANGLE, line_w=0.75)

    y = 1.78
    for head, color, items, h in sections:
        box(slide, 6.95, y + 0.04, 0.08, 0.26, fill=color, line=None,
            shape=MSO_SHAPE.RECTANGLE)
        label(slide, 7.1, y, 5.8, 0.34, head, size=13, bold=True, color=color,
              align=PP_ALIGN.LEFT)
        tb = slide.shapes.add_textbox(Inches(6.95), Inches(y + 0.36), Inches(5.98),
                                      Inches(h - 0.36))
        tf = tb.text_frame
        tf.margin_left = Inches(0.05)
        tf.margin_top = Inches(0.0)
        lines = []
        for lead, rest in items:
            lines.append([("■ ", False, color), (lead, True, DARK), (rest, False, DARK)])
        write(tf, lines, size=10.5, spacing=2)
        y += h

    if footer:
        fb = box(slide, 0.4, 6.98, 12.53, 0.4, fill=NAVY, line=None,
                 shape=MSO_SHAPE.RECTANGLE, align=PP_ALIGN.LEFT)
        fb.text_frame.margin_left = Inches(0.15)
        write(fb.text_frame, [[("共同结论：", True, ORANGE), (footer, False, WHITE)]],
              size=10.5)
    else:
        label(slide, 0.4, 7.05, 8, 0.3, "大模型发展趋势与网络诉求", size=9, color=GRAY,
              align=PP_ALIGN.LEFT)
        label(slide, 11.9, 7.05, 1.03, 0.3, f"{idx} / 3", size=9, color=GRAY,
              align=PP_ALIGN.RIGHT)

    slide.notes_slide.notes_text_frame.text = notes
    return slide


def panel_title(slide, text):
    label(slide, 0.55, 1.85, 6.0, 0.3, text, size=11, bold=True, color=NAVY,
          align=PP_ALIGN.LEFT)


# ---------------------------------------------------------------- slide 1

def diagram_inference(s):
    panel_title(s, "示意图：PD 分离 + 专家并行 + KV cache 多级池化")

    box(s, 0.75, 2.25, 1.4, 0.45, "用户请求", fill=WHITE, line=NAVY, size=10, bold=True)
    box(s, 2.75, 2.25, 2.4, 0.45, "全局调度 / 路由", fill=NAVY, line=None, size=10,
        bold=True, color=WHITE)
    arrow(s, 2.15, 2.475, 2.75, 2.475, color=NAVY)

    box(s, 0.75, 3.05, 2.45, 1.45, "", fill=WHITE, line=BLUE, line_w=1.25)
    label(s, 0.75, 3.1, 2.45, 0.3, "Prefill 池（算力密集）", size=9.5, bold=True, color=BLUE)
    gpu_grid(s, 1.2, 3.55, 4, 2, size=0.3, gap=0.12)

    box(s, 4.05, 3.05, 2.45, 1.45, "", fill=WHITE, line=BLUE, line_w=1.25)
    label(s, 4.05, 3.1, 2.45, 0.3, "Decode 池（带宽密集）", size=9.5, bold=True, color=BLUE)
    gpu_grid(s, 4.5, 3.55, 4, 2, size=0.3, gap=0.12)

    arrow(s, 3.6, 2.7, 1.975, 3.05, color=NAVY, width=1.25)
    arrow(s, 4.3, 2.7, 5.275, 3.05, color=NAVY, width=1.25)

    arrow(s, 3.2, 3.95, 4.05, 3.95, color=ORANGE, width=4)
    label(s, 3.12, 3.45, 1.0, 0.45, "KV 迁移\nRDMA", size=8.5, bold=True, color=ORANGE)
    label(s, 3.12, 4.02, 1.0, 0.3, "计入 TTFT", size=7.5, color=ORANGE)

    # long context growth
    box(s, 0.75, 4.7, 2.45, 1.2, "", fill=WHITE, line=MID_GRAY)
    label(s, 0.8, 4.73, 2.35, 0.28, "推理链 / 长上下文", size=9, bold=True, color=DARK)
    for i, (hh, t) in enumerate([(0.18, "MB"), (0.38, ""), (0.6, "GB")]):
        box(s, 1.05 + i * 0.35, 5.8 - hh, 0.25, hh, "", fill=ORANGE if i == 2 else GPU_FILL,
            line=None, shape=MSO_SHAPE.RECTANGLE)
        if t:
            label(s, 0.98 + i * 0.35, 5.8 - hh - 0.2, 0.4, 0.18, t, size=7, color=GRAY)
    label(s, 2.1, 5.05, 1.1, 0.75, "单请求 KV\n持续膨胀\nHBM 放不下", size=8, color=DARK,
          align=PP_ALIGN.LEFT)

    # EP all-to-all
    box(s, 4.05, 4.7, 2.45, 1.2, "", fill=WHITE, line=PURPLE, dash=True)
    label(s, 4.1, 4.73, 2.35, 0.28, "宽专家并行 all-to-all", size=9, bold=True, color=PURPLE)
    pts = [(4.45, 5.1), (5.1, 5.1), (4.45, 5.55), (5.1, 5.55)]
    r = 0.17
    for i in range(4):
        for j in range(i + 1, 4):
            arrow(s, pts[i][0], pts[i][1], pts[j][0], pts[j][1], color=PURPLE,
                  width=1.0, head=False)
    for k, (cx, cy) in enumerate(pts):
        e = box(s, cx - r, cy - r, 2 * r, 2 * r, f"E{k+1}", fill=PURPLE, line=None,
                size=7, color=WHITE, bold=True, shape=MSO_SHAPE.OVAL)
        e.text_frame.margin_left = e.text_frame.margin_right = 0
        e.text_frame.word_wrap = False
    label(s, 5.4, 5.05, 1.1, 0.75, "每层两次\n小消息高频\n尾时延→TPOT", size=8,
          color=DARK, align=PP_ALIGN.LEFT)

    # tiers
    label(s, 0.75, 6.0, 5.75, 0.25, "KV cache 多级池化（前缀命中后取回，按 block 离散读写）",
          size=9, bold=True, color=GREEN, align=PP_ALIGN.LEFT)
    tiers = ["HBM", "本地 DRAM", "SSD", "远端内存池"]
    tw, gap = 1.16, 0.37
    for i, t in enumerate(tiers):
        x = 0.75 + i * (tw + gap)
        box(s, x, 6.28, tw, 0.45, t, fill=WHITE if i else GPU_FILL, line=GREEN, size=9,
            bold=True, line_w=1.25)
        if i:
            arrow(s, x - gap + 0.03, 6.505, x - 0.03, 6.505, color=GREEN, width=1.75,
                  tail=True)


# ---------------------------------------------------------------- slide 2

def diagram_training(s):
    panel_title(s, "示意图：RL 后训练循环（训推共置）与跨域训练")

    box(s, 0.75, 2.25, 2.45, 2.0, "", fill=WHITE, line=BLUE, line_w=1.25)
    label(s, 0.75, 2.3, 2.45, 0.3, "训练集群", size=10, bold=True, color=BLUE)
    box(s, 0.95, 2.65, 2.05, 0.95, "", fill=LIGHT_BLUE, line=BLUE, dash=True)
    label(s, 0.95, 2.66, 2.05, 0.22, "超节点 scale-up 域（NVLink）", size=7.5, color=BLUE)
    gpu_grid(s, 1.12, 2.92, 5, 2, size=0.26, gap=0.1)
    label(s, 0.8, 3.65, 2.35, 0.55, "all-reduce / all-gather\nMoE all-to-all", size=8,
          color=DARK)

    box(s, 4.05, 2.25, 2.45, 2.0, "", fill=WHITE, line=TEAL, line_w=1.25)
    label(s, 4.05, 2.3, 2.45, 0.3, "Rollout 集群（推理副本 ×N）", size=10, bold=True,
          color=TEAL)
    for r_ in range(3):
        for c_ in range(4):
            box(s, 4.3 + c_ * 0.52, 2.72 + r_ * 0.45, 0.42, 0.34, "副本", fill=WHITE,
                line=TEAL, size=6.5, shape=MSO_SHAPE.RECTANGLE, line_w=0.75)
    label(s, 4.1, 4.02, 2.35, 0.22, "生成样本，常占大部分时间", size=7.5, color=DARK)

    arrow(s, 3.2, 2.8, 4.05, 2.8, color=ORANGE, width=3.5)
    label(s, 3.12, 2.33, 1.0, 0.45, "① 权重同步\n一对多广播", size=7.5, bold=True,
          color=ORANGE)
    arrow(s, 4.05, 3.75, 3.2, 3.75, color=GREEN, width=3.5)
    label(s, 3.12, 3.82, 1.0, 0.45, "② 轨迹 / logits\n/ 专家路由", size=7.5, bold=True,
          color=GREEN)

    box(s, 0.75, 4.55, 5.75, 0.55, "共享 DRAM / HBM 内存池：权重 · 样本 · 路由数据按需读写",
        fill=WHITE, line=GREEN, size=9, bold=True, color=GREEN, line_w=1.25)
    arrow(s, 1.975, 4.55, 1.975, 4.25, color=GREEN, width=1.5, tail=True)
    arrow(s, 5.275, 4.55, 5.275, 4.25, color=GREEN, width=1.5, tail=True)

    label(s, 0.75, 5.3, 5.75, 0.25, "跨数据中心训练（单园区电力受限）", size=9, bold=True,
          color=NAVY, align=PP_ALIGN.LEFT)
    for i, name in enumerate(["园区 A", "园区 B"]):
        x = 0.9 + i * 3.95
        box(s, x, 5.65, 1.5, 0.95, "", fill=WHITE, line=NAVY)
        label(s, x, 5.68, 1.5, 0.25, name, size=8.5, bold=True, color=NAVY)
        gpu_grid(s, x + 0.2, 5.98, 4, 1, size=0.22, gap=0.07)
        label(s, x, 6.27, 1.5, 0.25, "集群内：微秒级", size=7, color=GRAY)
    arrow(s, 2.4, 6.12, 4.85, 6.12, color=NAVY, width=2.25, dash=True, tail=True)
    label(s, 2.45, 5.72, 2.35, 0.35, "跨域链路", size=8.5, bold=True, color=NAVY)
    label(s, 2.45, 6.2, 2.35, 0.4, "带宽低一个量级以上\n时延毫秒级", size=7.5, color=DARK)


# ---------------------------------------------------------------- slide 3

def diagram_agentic(s):
    panel_title(s, "示意图：Agent 会话生命周期与 KV 换入换出")

    names = [("规划", 1.2), ("执行", 2.55), ("校验", 3.9)]
    for n, x in names:
        box(s, x - 0.38, 2.28, 0.76, 0.5, n, fill=NAVY, line=None, size=9, bold=True,
            color=WHITE, shape=MSO_SHAPE.OVAL)
    arrow(s, 1.58, 2.53, 2.17, 2.53, color=NAVY, width=1.25, tail=True)
    arrow(s, 2.93, 2.53, 3.52, 2.53, color=NAVY, width=1.25, tail=True)
    box(s, 4.85, 2.25, 1.65, 0.56, "工具 / 沙箱\n检索 / 浏览器", fill=WHITE, line=GRAY,
        size=8, bold=True)
    arrow(s, 4.28, 2.53, 4.85, 2.53, color=GRAY, width=1.25, tail=True)
    label(s, 0.75, 2.85, 5.75, 0.25, "多智能体协作：东西向小消息 / RPC，突发、不规则",
          size=8, color=DARK)

    label(s, 0.75, 3.18, 3.0, 0.25, "单个会话时间线（数十轮调用）", size=9, bold=True,
          color=NAVY, align=PP_ALIGN.LEFT)
    segs = [("第 1 轮\n推理", 1.0, BLUE), ("工具调用", 0.95, MID_GRAY),
            ("第 2 轮\n推理", 1.0, BLUE), ("工具调用", 0.95, MID_GRAY),
            ("第 N 轮\n推理", 1.0, BLUE)]
    x = 0.75
    centers = []
    for t, w, c in segs:
        box(s, x, 3.5, w - 0.05, 0.55, t, fill=c, line=None, size=8, bold=True,
            color=WHITE if c == BLUE else DARK, shape=MSO_SHAPE.RECTANGLE)
        centers.append(x + (w - 0.05) / 2)
        x += w + 0.12
    arrow(s, 0.75, 4.15, 6.5, 4.15, color=GRAY, width=1.25)
    label(s, 5.9, 4.17, 0.6, 0.2, "时间", size=7, color=GRAY)

    box(s, 0.75, 4.55, 5.75, 0.5, "", fill=GPU_FILL, line=BLUE, shape=MSO_SHAPE.RECTANGLE)
    label(s, 0.8, 4.55, 1.2, 0.5, "GPU HBM", size=9, bold=True, color=NAVY,
          align=PP_ALIGN.LEFT)
    box(s, 0.75, 5.55, 5.75, 0.5, "", fill=WHITE, line=GREEN, shape=MSO_SHAPE.RECTANGLE,
        line_w=1.25)
    label(s, 0.8, 5.55, 2.2, 0.5, "DRAM / 远端 KV 池", size=9, bold=True, color=GREEN,
          align=PP_ALIGN.LEFT)

    for i in (1, 3):
        cx = centers[i]
        arrow(s, cx - 0.2, 4.6, cx - 0.2, 5.55, color=ORANGE, width=2.25)
        label(s, cx - 0.62, 5.1, 0.42, 0.4, "换出", size=7.5, bold=True, color=ORANGE)
        nx = centers[i + 1] - 0.25
        arrow(s, nx, 5.55, nx, 4.6, color=GREEN, width=2.25)
        label(s, nx + 0.02, 5.08, 0.62, 0.45, "换入\n前缀复用", size=7, bold=True,
              color=GREEN, align=PP_ALIGN.LEFT)

    label(s, 0.75, 6.18, 3.0, 0.25, "会话迁移：上下文跨节点低成本搬移", size=9, bold=True,
          color=NAVY, align=PP_ALIGN.LEFT)
    box(s, 0.95, 6.45, 1.4, 0.36, "节点 A", fill=WHITE, line=NAVY, size=8.5, bold=True)
    box(s, 4.9, 6.45, 1.4, 0.36, "节点 B", fill=WHITE, line=NAVY, size=8.5, bold=True)
    arrow(s, 2.35, 6.63, 4.9, 6.63, color=NAVY, width=2.0)
    label(s, 2.4, 6.4, 2.45, 0.22, "会话 KV + 状态", size=7.5, color=DARK)


# ---------------------------------------------------------------- content

S1_NOTES = """核心判断：推理的瓶颈正在从 GPU 算力转向 KV cache 的存储与搬运，网络从“请求入口”变成推理系统内部的“内存总线”。

最显著的变化
- 推理型模型（test-time compute）：单个请求输出几千到几万 token，decode 时间长，KV cache 持续膨胀。
- 长上下文成为默认：128K 到百万级上下文，单请求 KV cache 可达 GB 级，HBM 放不下。
- Prefill / Decode 分离：两阶段放在不同 GPU 池，KV cache 必须跨机传递。
- 大规模 MoE + 宽专家并行：专家分散到几十到上百张卡，每层都要 dispatch / combine。
- KV cache 池化：HBM → DRAM → SSD → 远端内存池多级存放，前缀命中后再取回。

流量特征
- PD 之间：大块、突发的 KV cache 迁移，时延直接计入 TTFT。
- 专家并行：每层两次 all-to-all，消息小、频率高，尾时延决定 TPOT。
- KV 池：DRAM / HBM 之间高带宽读写，按 block 索引的离散读取。

对网络的诉求
1. 内存级带宽：单机多网卡聚合，逼近 PCIe / NIC 物理上限；HBM↔DRAM↔远端 DRAM 全程零拷贝。
2. 微秒级尾时延：GPU 直接发起通信、少 CPU 介入，拥塞时 P99 可控。
3. 拓扑感知：GPU、NIC、NUMA、NVLink 亲和关系进入调度；不在同一 PCIe 的网卡可借道 NVLink。
4. 离散访问高效：scatter / gather 原语，大量小 block 一次性批量搬运。
5. SLO 隔离：KV 迁移、专家通信、存储回写共享网卡时，时延敏感流量优先。"""

S2_NOTES = """核心判断：训练算力增量越来越多投向强化学习后训练，训练集群里混入了推理负载，流量从规整的集合通信变成“集合通信 + 权重同步 + 样本流转”的混合形态。

最显著的变化
- RL 后训练规模化：rollout 往往占据大部分时间，训练与推理在同一任务里交替。
- 训推共置 / 异步 RL：每轮训练后把新权重下发到 rollout 集群，副本可能多达数百个。
- MoE 成为主流：训练中也有 all-to-all，随专家数和并行度增长。
- Scale-up 域扩大：超节点把 NVLink 类互联扩展到数十到上百卡，scale-up 与 scale-out 分层。
- 更低精度（FP8 / FP4）：计算变快而通信量下降有限，通信占比被动升高。
- 跨数据中心训练：单园区电力受限，训练开始跨楼、跨园区甚至跨地域。

流量特征
- 权重同步：训练侧到 rollout 侧的一对多大块广播，频繁、周期性，影响 GPU 空闲时间。
- 样本流转：轨迹、logits、专家路由信息回流训练侧，数据量大且有时效性。
- 集合通信：all-reduce / all-gather / all-to-all 仍是主体，但与上述流量共享一张网。
- 跨域链路：带宽比集群内低一个量级以上，时延从微秒变成毫秒。

对网络的诉求
1. 高效一对多：权重广播流水线化、多轨并行，避免所有副本同时从同一源拉取。
2. 训推一体的内存池：权重、样本、路由数据放入共享 DRAM / HBM 池，按需读写。
3. Scale-up 与 scale-out 协同：节点内走 NVLink，节点间多网卡负载均衡。
4. 万卡级可靠性：链路、网卡故障秒级隔离与自动恢复，任务不整体重启。
5. 跨域友好：长距离拥塞控制、流量压缩、计算与通信重叠。"""

S3_NOTES = """核心判断：Agent 把一次调用变成持续几十分钟到数小时的会话，上下文需要长期保存和反复复用，负载从“吞吐型”转向“状态型 + 交互型”。

最显著的变化
- 多轮、长程任务：一个任务包含几十到上百轮模型调用，上下文持续累积。
- 工具调用与环境交互：推理与代码执行、检索、浏览器、沙箱交替，GPU 需在等待期间换出会话。
- 多智能体协作：规划、执行、校验等角色并行，彼此传递上下文和中间结果。
- 记忆分层：短期上下文（KV cache）、会话记忆、长期知识库并存。
- 并发会话数激增：单用户可同时运行多个 Agent，后台任务不受人类交互速度约束。

流量特征
- 会话级 KV 复用：每轮只追加少量 token，历史前缀反复命中；暂停时换出，恢复时快速取回。
- 换入换出频繁：工具调用期间释放 HBM，返回后重新加载，大量中等块双向搬运。
- 东西向流量增加：Agent 之间、Agent 与工具服务之间大量小消息和 RPC，突发、不规则。
- 粘性与迁移并存：会话倾向留在原节点复用缓存，负载均衡又要求能迁移。

对网络的诉求
1. 分布式 KV 内存池：跨节点共享的 KV cache 存储，按前缀索引定位，命中后直接 RDMA 读取。
2. 快速换入换出：HBM 与 DRAM、远端内存之间高带宽搬运，恢复时间远小于重算。
3. 会话可迁移：上下文低成本跨节点搬移，调度不被缓存位置“绑死”。
4. 混合流量 QoS：区分交互时延敏感流量、后台批量搬运和存储流量。
5. 规模化连接管理：大量并发会话的连接建立、元数据交换要轻量，避免控制面成为瓶颈。

共同结论：推理、训练、Agent 都在把“状态”（KV cache、权重、样本、会话上下文）推到网络上。网络关键指标从单纯的集合通信带宽，扩展为内存池级带宽、尾时延、拓扑感知和多类流量隔离。"""


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    s1 = frame(
        prs, 1, "推理：从“算一次”到“存得住、调得动”的状态化服务",
        "推理瓶颈从 GPU 算力转向 KV cache 的存储与搬运，网络从“请求入口”变成推理系统内部的“内存总线”。",
        [
            ("最显著的变化", BLUE, [
                ("推理型模型：", "单请求输出数千至数万 token，KV 持续膨胀"),
                ("长上下文成为默认：", "128K 到百万级，单请求 KV 达 GB 级"),
                ("PD 分离：", "prefill / decode 分池部署，KV 必须跨机迁移"),
                ("大规模 MoE + 宽专家并行：", "每层都要 dispatch / combine"),
                ("KV cache 池化：", "HBM → DRAM → SSD → 远端内存池多级存放"),
            ], 1.95),
            ("流量特征", ORANGE, [
                ("PD 之间：", "大块突发 KV 迁移，时延直接计入 TTFT"),
                ("专家并行：", "高频小消息 all-to-all，尾时延决定 TPOT"),
                ("KV 池：", "高带宽、按 block 索引的离散读写"),
            ], 1.25),
            ("对网络的诉求", GREEN, [
                ("内存级带宽：", "多网卡聚合逼近物理上限，全路径零拷贝"),
                ("微秒级尾时延：", "GPU 直发通信、少 CPU 介入，P99 可控"),
                ("拓扑感知：", "GPU / NIC / NUMA / NVLink 亲和进入调度"),
                ("离散访问高效：", "scatter / gather 批量搬运大量小 block"),
                ("SLO 隔离：", "共享网卡时时延敏感流量优先"),
            ], 1.92),
        ],
        S1_NOTES,
    )
    diagram_inference(s1)

    s2 = frame(
        prs, 2, "训练：从“预训练单一范式”到“预训练 + 大规模强化学习”",
        "算力增量投向 RL 后训练，集群混入推理负载，流量变成“集合通信 + 权重同步 + 样本流转”的混合形态。",
        [
            ("最显著的变化", BLUE, [
                ("RL 后训练规模化：", "rollout 占大部分时间，训推交替运行"),
                ("训推共置 / 异步 RL：", "每轮后向数百副本下发新权重"),
                ("MoE 成为主流：", "训练中 all-to-all 随专家数增长"),
                ("Scale-up 域扩大：", "超节点数十到上百卡，两层互联分层"),
                ("更低精度 FP8 / FP4：", "计算变快，通信占比被动升高"),
                ("跨数据中心训练：", "单园区电力受限，跨楼、跨园区、跨地域"),
            ], 2.18),
            ("流量特征", ORANGE, [
                ("权重同步：", "一对多大块广播，周期性，影响 GPU 空闲"),
                ("样本流转：", "轨迹 / logits / 路由回流，量大且有时效"),
                ("跨域链路：", "带宽低一个量级以上，时延毫秒级"),
            ], 1.2),
            ("对网络的诉求", GREEN, [
                ("高效一对多：", "权重广播流水线化、多轨并行"),
                ("训推一体内存池：", "权重、样本、路由数据共享池按需读写"),
                ("两层协同：", "节点内走 NVLink，节点间多网卡均衡"),
                ("万卡级可靠性：", "故障秒级隔离与自动恢复"),
                ("跨域友好：", "长距拥塞控制、压缩、计算通信重叠"),
            ], 1.74),
        ],
        S2_NOTES,
    )
    diagram_training(s2)

    s3 = frame(
        prs, 3, "Agentic AI：从“一问一答”到“长程、多轮、多体协作”",
        "Agent 把一次调用变成持续数十分钟到数小时的会话，负载从“吞吐型”转向“状态型 + 交互型”。",
        [
            ("最显著的变化", BLUE, [
                ("多轮、长程任务：", "数十到上百轮调用，上下文持续累积"),
                ("工具调用与环境交互：", "等待期间 GPU 需换出会话"),
                ("多智能体协作：", "规划、执行、校验并行，互传上下文"),
                ("记忆分层：", "KV cache、会话记忆、长期知识库并存"),
                ("并发会话激增：", "后台任务不受人类交互速度约束"),
            ], 1.95),
            ("流量特征", ORANGE, [
                ("会话级 KV 复用：", "每轮少量追加，历史前缀反复命中"),
                ("换入换出频繁：", "工具调用期间释放 HBM，返回后重载"),
                ("东西向小消息：", "Agent 与工具间 RPC，突发、不规则"),
            ], 1.25),
            ("对网络的诉求", GREEN, [
                ("分布式 KV 内存池：", "按前缀索引定位，命中后 RDMA 直读"),
                ("快速换入换出：", "恢复时间远小于重算时间"),
                ("会话可迁移：", "调度不被缓存位置“绑死”"),
                ("混合流量 QoS：", "交互、批量搬运、存储流量分级"),
                ("轻量连接管理：", "海量会话下控制面不成为瓶颈"),
            ], 1.92),
        ],
        S3_NOTES,
        footer="三个方向都在把“状态”推到网络上，关键指标从集合通信带宽扩展为内存池级带宽、尾时延、拓扑感知与多类流量隔离。",
    )
    diagram_agentic(s3)

    out_dir = os.path.dirname(os.path.abspath(__file__))
    prs.save(os.path.join(out_dir, "llm-trends-network-requirements.pptx"))


if __name__ == "__main__":
    main()
